#!/usr/bin/env python3
"""Read a workbook without changing it; emit private, deterministic import data."""
import argparse
import datetime as dt
import hashlib
import json
import re
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import openpyxl

MONTH = re.compile(r'^\d{4}_\d{2}$')


def serial(value):
    if isinstance(value, openpyxl.worksheet.formula.ArrayFormula):
        return {'formulaType': 'array', 'reference': value.ref, 'text': value.text}
    if isinstance(value, openpyxl.worksheet.formula.DataTableFormula):
        return {'formulaType': 'dataTable', **dict(value)}
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    return value


def minor(value):
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        raise ValueError('Amount must be numeric')
    return int((Decimal(str(value)) * 100).quantize(Decimal('1'), rounding=ROUND_HALF_UP))


def stable_id(sheet, reference, kind):
    return hashlib.sha256(f'wallet-v1|{sheet}|{reference}|{kind}'.encode()).hexdigest()[:32]


def cell_record(cell, cached):
    result = {'reference': cell.coordinate, 'value': serial(cell.value),
              'cachedValue': serial(cached.value), 'dataType': cell.data_type,
              'numberFormat': cell.number_format}
    if cell.data_type == 'f':
        result['formula'] = serial(cell.value)
    if cell.comment:
        result['comment'] = {'text': cell.comment.text, 'author': cell.comment.author}
    if cell.hyperlink:
        result['hyperlink'] = {'target': cell.hyperlink.target, 'location': cell.hyperlink.location}
    return result


def populated(sheet):
    # Avoid materializing the huge blank styled ranges used by full-column formulas.
    return sorted((c for c in sheet._cells.values() if c.value is not None or c.comment),
                  key=lambda c: (c.row, c.column))


def migrate(source):
    source = Path(source)
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    formulas = openpyxl.load_workbook(source, data_only=False)
    cached = openpyxl.load_workbook(source, data_only=True)
    result = {'schemaVersion': 1, 'source': {'filename': source.name, 'sha256': digest},
              'transactions': [], 'categories': [], 'modules': [], 'historicalSummaries': [],
              'reviewItems': [], 'sourceArchive': {'sheets': []}}
    reconciliations = []
    category_set = set()

    def review(sheet, reference, reason, **extra):
        result['reviewItems'].append({'id': stable_id(sheet, reference, reason),
            'sourceSheet': sheet, 'sourceReference': reference, 'reason': reason, **extra})

    def module(sheet, row, kind, cells, **fields):
        result['modules'].append({'id': stable_id(sheet, str(row), kind), 'type': kind,
            'title': fields.pop('title', f'{sheet} · row {row}'), 'sourceSheet': sheet,
            'fields': {'sourceReference': f'row {row}', 'cells': cells, **fields}})

    def row_fields(sheet, row, kind):
        """Add conveniences using explicit source columns; original cells remain authoritative."""
        row = int(row)
        ws = cached[sheet]
        def v(col):
            return serial(ws[f'{col}{row}'].value)
        def money(col):
            value = ws[f'{col}{row}'].value
            return minor(value) if isinstance(value,(int,float)) and not isinstance(value,bool) else None
        if kind == 'debt':
            label = v('C')
            entries = []
            if 4 <= row <= 28:
                for amountcol,statuscol in [('D','H'),('E','I'),('F','J')]:
                    amount = money(amountcol)
                    if amount is None:
                        continue
                    status = v(statuscol)
                    paid = amount if str(status).upper() == 'Y' else 0 if str(status).upper() == 'N' else None
                    entries.append({'person':ws[f'{amountcol}3'].value,'amountMinor':amount,
                        'paidStatus':status,'paidMinor':paid,
                        'remainingMinor':None if paid is None else amount-paid,
                        'amountReference':f'{amountcol}{row}','statusReference':f'{statuscol}{row}'})
            return {'description':label,'entries':entries,'isSummary':row in (29,31,32),
                    'currency':'HKD'}
        if kind == 'card':
            return {'bank':v('U'),'card':v('V'),'annualFeeMinor':money('W'),
                    'usage':v('X'),'amountMinor':money('Y'),'status':v('Z'),'notes':v('AA'),
                    'isRecord':row >= 6 and bool(v('U')) and bool(v('V')) and v('U') != 'Type',
                    'currency':'HKD'}
        if kind == 'insurance':
            valid = bool(v('U')) and bool(v('V')) and v('U') not in ('Type','Insurance')
            return {'policyType':v('U'),'policyName':v('V'),'paymentMode':v('W'),
                    'premiumMinor':money('X'),'premiumSourceLabel':'Premium (Y)',
                    'remarks':v('Y'),'isRecord':valid,'currency':'HKD',
                    'premiumInterpretation':'Preserved amount and payment mode; do not infer annualization.'}
        if kind == 'tax':
            return {'year':v('C'),'expectedAnnualIncomeMinor':money('D'),
                    'projectedTaxMinor':money('E'),'monthlyTaxSavingMinor':money('F'),
                    'isRecord':isinstance(ws[f'C{row}'].value,(int,float)),
                    'currency':'HKD','estimateSource':serial(ws['E2'].value)}
        if kind == 'investment':
            entries=[]
            for datecol,directioncol,amountcol,strategy in [('B','C','D','QQQ'),('J','K','L','Tax saving USD cash'),('P','Q','R','Monthly saving HKD cash')]:
                if isinstance(ws[f'{datecol}{row}'].value,(dt.date,dt.datetime)):
                    entry={'strategy':strategy,'date':v(datecol),'direction':v(directioncol),
                        'hkdAmountMinor':money(amountcol),'sourceReference':f'{datecol}{row}:{amountcol}{row}'}
                    if datecol=='B':entry.update({'exchangeRate':v('E'),'usdAmount':v('F'),'units':v('G'),'unitPrice':v('H')})
                    if datecol=='J':entry.update({'exchangeRate':v('M'),'usdAmount':v('N')})
                    entries.append(entry)
            return {'entries':entries,'isRecord':bool(entries)}
        if kind == 'project':
            datecol,amountcol=('D','E') if sheet=='Rent' else ('E','D')
            return {'item':v('C'),'date':v(datecol),'amountMinor':money(amountcol),
                    'task':v('I') if sheet=='Rent' else None,'progress':v('J') if sheet=='Rent' else None,
                    'deliveryDate':v('K') if sheet=='Rent' else v('F'),'currency':'HKD',
                    'isRecord':bool(v('C')) and money(amountcol) is not None}
        if kind == 'trip':
            if sheet=='KYUSHU':return {'category':v('D'),'item':v('E'),'originalCurrency':v('L'),
                'originalAmount':v('M'),'jpyAmount':v('F'),'hkdPersonalShareMinor':money('G'),
                'paymentMethod':v('H'),'payer':v('I'),'splitwiseStatus':v('J'),
                'exchangeRate':serial(ws['P1'].value),'isRecord':row>=4 and bool(v('E'))}
            if sheet=='JP202406':return {'date':v('D'),'category':v('E'),'item':v('F'),
                'include':v('G'),'originalCurrency':v('H'),'paymentMethod':v('I'),
                'jpyAmount':v('J'),'hkdAmountMinor':money('K'),'hkdFinalMinor':money('M'),
                'remarks':v('N'),'exchangeRate':serial(ws['B1'].value),'isRecord':row>=3 and bool(v('F'))}
            return {'item':v('E'),'description':v('F'),'hkdAmountMinor':money('G'),
                'twdAmount':v('H'),'exchangeRate':v('I'),'remarks':v('J'),
                'isRecord':row>=3 and money('G') is not None}
        return {}

    for sheet in formulas:
        cs = cached[sheet.title]
        archive_cells = [cell_record(c, cs[c.coordinate]) for c in populated(sheet)]
        result['sourceArchive']['sheets'].append({'name': sheet.title, 'state': sheet.sheet_state,
            'maxRow': sheet.max_row, 'maxColumn': sheet.max_column,
            'mergedRanges': [str(r) for r in sheet.merged_cells.ranges],
            'cells': archive_cells})
        rows = defaultdict(list)
        for c in archive_cells:
            rows[re.search(r'\d+', c['reference']).group()].append(c)

        if MONTH.match(sheet.title):
            period = sheet.title.replace('_', '-')
            for kind, dcol, labelcol, acol, summary in [
                    ('expense', 'B', 'C', 'D', 'M5'), ('income', 'G', 'H', 'I', 'M4')]:
                exact_total = Decimal(0)
                rounded_total = 0
                unresolved = 0
                for row in range(6, sheet.max_row + 1):
                    ref = f'{acol}{row}'
                    original = sheet[ref].value
                    label = sheet[f'{labelcol}{row}'].value
                    day = cs[f'{dcol}{row}'].value
                    if original is None:
                        if label is not None or day is not None:
                            review(sheet.title, ref, 'Blank amount; preserved as template or incomplete entry')
                        continue
                    amount = cs[ref].value
                    if not isinstance(amount, (int, float)) or isinstance(amount, bool):
                        review(sheet.title, ref, 'Amount has no numeric cached result', originalValue=serial(original))
                        unresolved += 1
                        continue
                    money = minor(amount)
                    exact_total += Decimal(str(amount))
                    rounded_total += money
                    if isinstance(day, (dt.date, dt.datetime)):
                        date = day.strftime('%Y-%m-%d')
                    else:
                        date = None
                        review(sheet.title, f'{dcol}{row}', 'Missing or non-date transaction date', originalValue=serial(day))
                    if not label:
                        review(sheet.title, f'{labelcol}{row}', 'Missing category')
                    original_label = str(label) if label is not None else 'Uncategorized'
                    category_set.add(original_label)
                    notes = []
                    if date and date[:7] != period:
                        review(sheet.title, f'{dcol}{row}', 'Date differs from reporting month', date=date, reportingPeriod=period)
                    for c in (sheet[f'{dcol}{row}'], sheet[f'{labelcol}{row}'], sheet[ref]):
                        if c.comment:
                            notes.append(c.comment.text)
                    transaction = {'id': stable_id(sheet.title, ref, kind), 'date': date,
                        'reportingPeriod': period, 'kind': kind, 'category': original_label,
                        'originalCategory': original_label, 'description': original_label,
                        'account': 'Historical / unspecified', 'currency': 'HKD',
                        'amountMinor': money, 'personalShareMinor': money,
                        'sourceSheet': sheet.title, 'sourceReference': f'{dcol}{row}:{acol}{row}'}
                    if notes:
                        transaction['notes'] = '\n'.join(notes)
                    result['transactions'].append(transaction)
                expected = cs[summary].value
                expected_minor = minor(expected) if isinstance(expected, (int, float)) else None
                reconciliation = {'sourceSheet': sheet.title, 'kind': kind,
                    'sourceReference': summary, 'rawTotal': str(exact_total),
                    'importedMinor': rounded_total, 'cachedMinor': expected_minor,
                    'differenceMinor': None if expected_minor is None else rounded_total - expected_minor,
                    'rawDifferenceMinor': None if expected_minor is None else minor(exact_total) - expected_minor,
                    'unresolvedAmounts': unresolved,
                    'status': 'unavailable' if expected_minor is None else
                        ('match' if minor(exact_total) == expected_minor and unresolved == 0 else 'difference')}
                reconciliations.append(reconciliation)
                if reconciliation['status'] != 'match':
                    review(sheet.title, summary, 'Monthly total needs reconciliation', reconciliation=reconciliation)
            # Side tables are snapshots, not ledger expenses. All populated side cells survive.
            section = 'card'
            for row in sorted(rows, key=int):
                side = [c for c in rows[row] if openpyxl.utils.cell.coordinate_from_string(c['reference'])[0] not in
                        ('A','B','C','D','E','F','G','H','I','J','K','L','M','N','O','P','Q','R','S','T')]
                if not side:
                    continue
                if any(c['value'] == 'Insurance' for c in side):
                    section = 'insurance'
                module(sheet.title, row, section, side, **row_fields(sheet.title,row,section), reportingPeriod=period,
                       title=f'{section.title()} · {period} · row {row}')
            continue

        if sheet.title in ('Summary->', 'Record->'):
            continue
        if sheet.title == 'Saving':
            for col in range(4, 10):
                year = str(sheet.cell(2, col).value)
                for row in (7, 8, 9, 43, 55, 57):
                    cell = sheet.cell(row, col)
                    result['historicalSummaries'].append({'id': stable_id('Saving', cell.coordinate, 'summary'),
                        'year': year, 'label': sheet.cell(row, 3).value,
                        'sourceSheet': 'Saving', 'sourceReference': cell.coordinate,
                        'value': serial(cs[cell.coordinate].value), 'formula': cell.value})
                for row in list(range(12, 43)) + list(range(46, 55)):
                    cell = sheet.cell(row, col)
                    val = cs[cell.coordinate].value
                    label = sheet.cell(row, 3).value
                    if cell.value is None or val in (None, 0):
                        continue
                    # Saving!H/I43 excludes rows 28:31 because their parent row27 is already included.
                    if row in range(28, 32) and col in (8, 9):
                        continue
                    if not isinstance(val, (int, float)):
                        review('Saving', cell.coordinate, 'Non-numeric exceptional item preserved', originalValue=serial(cell.value))
                        continue
                    category_set.add(str(label))
                    money = minor(val) * (-1 if row < 43 else 1)
                    result['transactions'].append({'id': stable_id('Saving', cell.coordinate, 'adjustment'),
                        'date': None, 'reportingPeriod': year, 'kind': 'adjustment',
                        'category': str(label), 'originalCategory': str(label), 'description': str(label),
                        'account': 'Historical annual adjustment', 'currency': 'HKD',
                        'amountMinor': money, 'personalShareMinor': money,
                        'sourceSheet': 'Saving', 'sourceReference': cell.coordinate,
                        'notes': 'Undated annual adjustment; auxiliary detail sheets excluded from ledger to avoid duplication.'})
            # Keep the complete table for auditing excluded subrows and labels.
        kind = {'Debt': 'debt', 'Rent': 'project', 'Coffee': 'project', 'KYUSHU': 'trip',
                'JP202406': 'trip', 'TW': 'trip', 'TaxProjection': 'tax',
                'MonthlyInvestment': 'investment', 'Saving': 'saving', 'Income': 'income-summary'}.get(sheet.title, 'expense-summary')
        for row in sorted(rows, key=int):
            cells = rows[row]
            titles = [str(c['cachedValue']) for c in cells if c['dataType'] != 'f' and
                      isinstance(c['cachedValue'], str)]
            module(sheet.title, row, kind, cells, **row_fields(sheet.title,row,kind),title=' · '.join(titles[:2]) or f'{sheet.title} · row {row}')
        if re.match(r'^\d{4}_Expense$', sheet.title):
            year = sheet.title[:4]
            total_ref = 'K3' if year == '2021' else 'P3'
            expected = cs[total_ref].value
            result['historicalSummaries'].append({'id': stable_id(sheet.title, total_ref, 'annual'),
                'year': year, 'label': 'Annual expense', 'value': serial(expected),
                'sourceSheet': sheet.title, 'sourceReference': total_ref})

    for summary in result['historicalSummaries']:
        if summary['label'] != 'Annual expense':
            continue
        imported = sum(t['amountMinor'] for t in result['transactions']
            if t['kind'] == 'expense' and t['reportingPeriod'].startswith(summary['year']))
        expected = summary['value']
        reconciliations.append({'sourceSheet': summary['sourceSheet'], 'kind': 'annual-expense',
            'sourceReference': summary['sourceReference'], 'importedMinor': imported,
            'cachedMinor': minor(expected) if isinstance(expected, (int, float)) else None,
            'differenceMinor': imported - minor(expected) if isinstance(expected, (int,float)) else None,
            'status': 'unavailable' if not isinstance(expected,(int,float)) else
                ('match' if imported == minor(expected) else 'difference')})

    # Evaluate savings from the imported ledger including explicitly signed annual adjustments.
    if 'Saving' in formulas.sheetnames:
        for col in range(4,10):
            year = str(formulas['Saving'].cell(2,col).value)
            totals = defaultdict(int)
            for t in result['transactions']:
                if t['reportingPeriod'].startswith(year):
                    totals[t['kind']] += t['amountMinor']
            imported = totals['income'] - totals['expense'] + totals['adjustment']
            ref = formulas['Saving'].cell(57,col).coordinate
            value = cached['Saving'][ref].value
            expected = minor(value) if isinstance(value,(int,float)) else None
            income_value=cached['Saving'].cell(7,col).value
            income_expected=minor(income_value) if isinstance(income_value,(int,float)) else None
            reconciliations.append({'sourceSheet':'Saving','sourceReference':formulas['Saving'].cell(7,col).coordinate,
                'kind':'annual-income','year':year,'importedMinor':totals['income'],
                'cachedMinor':income_expected,'differenceMinor':None if income_expected is None else totals['income']-income_expected,
                'status':'unavailable' if income_expected is None else ('match' if income_expected == totals['income'] else 'difference')})
            reconciliations.append({'sourceSheet':'Saving', 'sourceReference':ref, 'kind':'annual-saving',
                'year':year, 'importedMinor':imported, 'cachedMinor':expected,
                'differenceMinor': None if expected is None else imported - expected,
                'status':'unavailable' if expected is None else ('match' if expected == imported else 'difference')})
    result['categories'] = sorted(category_set, key=lambda s: (s.casefold(), s))
    report = {'schemaVersion': 1, 'source': result['source'], 'sheetCount':len(formulas.sheetnames),
        'monthlySheetCount':sum(bool(MONTH.match(n)) for n in formulas.sheetnames),
        'transactionCount': len(result['transactions']), 'moduleCount':len(result['modules']),
        'reviewItemCount':len(result['reviewItems']), 'reconciliations':reconciliations,
        'notes':['Cached formula values are preserved, not recalculated.',
            'All source cells and the original workbook remain available; auxiliary records are not additional ledger transactions.',
            'Each transaction rounds to cents; original precision is preserved in the source archive.',
            'Annual adjustments have no invented date; reportingPeriod is YYYY.',
            'Ambiguous duplicate purchases across monthly and exceptional annual records require review.']}
    # Every non-matching annual reconciliation must remain actionable in the mobile review list.
    for item in reconciliations:
        if item['kind'].startswith('annual') and item['status'] != 'match':
            review(item['sourceSheet'],item['sourceReference'],'Annual report differs or unavailable',reconciliation=item)
    report['reviewItemCount'] = len(result['reviewItems'])
    formulas.close()
    cached.close()
    if hashlib.sha256(source.read_bytes()).hexdigest() != digest:
        raise RuntimeError('Source changed during migration')
    return result, report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source', nargs='?', default='Expense and Saving.xlsx')
    parser.add_argument('--output-dir', default='private')
    args = parser.parse_args()
    result, report = migrate(args.source)
    directory = Path(args.output_dir)
    directory.mkdir(parents=True, exist_ok=True)
    for name, value in [('migration.json',result),('migration-report.json',report),
                        ('import.json',{k:v for k,v in result.items() if k != 'sourceArchive'})]:
        target = directory / name
        target.write_text(json.dumps(value,ensure_ascii=False,indent=2,allow_nan=False),encoding='utf-8')
        target.chmod(0o600)
    print(json.dumps({k:report[k] for k in ('sheetCount','monthlySheetCount','transactionCount','moduleCount','reviewItemCount')}))


if __name__ == '__main__':
    main()
