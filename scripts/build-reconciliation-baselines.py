"""Read-only workbook baselines; generated financial data belongs in private/ only."""
import argparse,json,re,hashlib
from pathlib import Path
from decimal import Decimal,ROUND_HALF_UP
import openpyxl

def minor(value):
    return int((Decimal(str(value))*100).quantize(Decimal('1'),rounding=ROUND_HALF_UP))
def main():
    parser=argparse.ArgumentParser();parser.add_argument('workbook');parser.add_argument('--import-file',default='private/import.json');parser.add_argument('--output',default='private/reconciliation-records.json');args=parser.parse_args()
    source=Path(args.workbook);before=hashlib.sha256(source.read_bytes()).hexdigest();data=json.loads(Path(args.import_file).read_text());assert data['source']['sha256']==before
    workbook=openpyxl.load_workbook(source,data_only=True,read_only=True);records=[]
    for sheet in workbook:
        if not re.fullmatch(r'\d{4}_\d{2}',sheet.title):continue
        period=sheet.title.replace('_','-');rows=[t for t in data['transactions'] if t['reportingPeriod']==period and t['kind'] in ('expense','income')]
        income,expense=sheet['M4'].value,sheet['M5'].value
        if not all(isinstance(n,(int,float)) for n in [income,expense]):continue
        categories=[]
        for cells in sheet.iter_rows(min_row=6,max_row=sheet.max_row,min_col=16,max_col=18):
            label,count,amount=(c.value for c in cells)
            if isinstance(label,str) and isinstance(amount,(int,float)):
                categories.append({'category':label,'workbookCount':int(count or 0),'workbookMinor':minor(amount),'sourceReference':f'P{cells[0].row}:R{cells[0].row}'})
        ident=f'{before}:reconciliation:{period}';payload={'id':ident,'type':'reconciliation','title':f'Workbook reconciliation {period}','sourceSheet':sheet.title,'fields':{'period':period,'incomeMinor':minor(income),'expenseMinor':minor(expense),'netMinor':minor(income)-minor(expense),'incomeCount':sum(t['kind']=='income' for t in rows),'expenseCount':sum(t['kind']=='expense' for t in rows),'refs':{'income':f'{sheet.title}!M4','expense':f'{sheet.title}!M5','net':f'{sheet.title}!M6'},'baselineTransactions':rows,'categoryChecks':categories}}
        records.append({'id':ident,'entity':'module','payload':payload,'version':0})
    workbook.close();assert hashlib.sha256(source.read_bytes()).hexdigest()==before
    target=Path(args.output);assert target.parts[0]=='private','Financial output must remain private';target.parent.mkdir(exist_ok=True,parents=True);target.write_text(json.dumps({'records':records},ensure_ascii=False,separators=(',',':')))
    print(json.dumps({'monthlyBaselines':len(records),'output':'private reconciliation payload','largestRecordBytes':max(len(json.dumps(r).encode()) for r in records)}))
if __name__=='__main__':main()
